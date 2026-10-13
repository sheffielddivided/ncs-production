"""
Importerer normpris-data fra Petroleumsprisrådet (xlsx-filer fra regjeringen.no)
og lagrer dem som kompakte JSON-filer under data/.

Rådata ligger i data/normpris-raw/. regjeringen.no ligger bak en Cloudflare-
utfordring og kan ikke hentes automatisk, så nye kvartalsfiler lastes ned
manuelt og legges i den mappa. Deretter kjører man:

    pip install openpyxl
    python fetch_normpris.py

Kjøres også automatisk av .github/workflows/update-normpris.yml når filer i
data/normpris-raw/ endres.

Hver fil har to ark vi bruker:
  'Differensialer'                  felt x måned, $/fat (D/F/H = måned 1/2/3)
  'Daglige normpriser (Dt. Brent)'  daglig Datert Brent, $/fat, tre
                                    dato/pris-par per rad (B/C, E/F, H/I)
Daglige normpriser per felt (NOK og $) og valutakurser leses ikke.

Noen kvartaler har i tillegg tilleggsfiler som retter eller fyller ut enkelte
felt/måneder ('...-troll-februar', '...-gudrun-juni' osv). Verdier fra en
tilleggsfil overstyrer hovedfilen, så en verdi som er rettet i ettertid vinner.
Manglende celle = ingen verdi (trolig ikke solgt olje den måneden); slike
måneder utelates, de lagres ikke som 0 eller null.

Utdata er deterministisk (sorterte rader og nøkler), slik at filene bare
endrer seg når dataene har endret seg.

  normpris_brent.json  {"columns":["date","usd"], "rows":[["2021-01-01",54.3],...]}
  normpris_diff.json   {"months":["2021-01",...], "fields":{"TROLL":{"2021-01":2.2,...}}}
"""

import glob
import json
import os
import re
import sys
from collections import defaultdict
from datetime import date, datetime

import openpyxl

RAW_DIR = os.path.join('data', 'normpris-raw')
OUT_DIR = 'data'

DIFF_SHEET  = 'Differensialer'
BRENT_SHEET = 'Daglige normpriser (Dt. Brent)'

MONTHS = {
    'januar': 1, 'februar': 2, 'mars': 3, 'april': 4, 'mai': 5, 'juni': 6,
    'juli': 7, 'august': 8, 'september': 9, 'oktober': 10, 'november': 11,
    'desember': 12,
}

# Kolonnene i Differensialer: D/F/H er måned 1/2/3 i kvartalet.
DIFF_COLS = (4, 6, 8)
# Brent-arket: (dato-kolonne, pris-kolonne) per måned.
BRENT_COLS = ((2, 3), (5, 6), (8, 9))

TITLE_RE = re.compile(r'(\d)Q\s*(\d{4})')


class ParseError(Exception):
    pass


def is_supplement(path):
    """Tilleggsfiler har et suffiks etter '-endelige-normpriser'."""
    name = os.path.basename(path).lower()
    return not re.fullmatch(r'np-\dq-\d{4}-endelige-normpriser\.xlsx', name)


def _quarter(ws, path):
    title = str(ws['B2'].value or '')
    m = TITLE_RE.search(title)
    if not m:
        raise ParseError(f"{path}: fant ikke kvartal/år i tittel {title!r}")
    return int(m.group(1)), int(m.group(2))


def _num(value, where):
    if value is None or value == '':
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ParseError(f"{where}: forventet tall, fant {value!r}")
    return float(value)


def parse_diffs(wb, path):
    """Returnerer liste av (felt, 'YYYY-MM', verdi) for én fil."""
    ws = wb[DIFF_SHEET]
    quarter, year = _quarter(ws, path)
    first_month = (quarter - 1) * 3 + 1

    # Overskriftene i rad 4 oppgir hvilke måneder kolonnene gjelder. Sjekk dem
    # mot kvartalet fra tittelen, så en forskjøvet kolonne ikke går ubemerket.
    month_of = {}
    for i, col in enumerate(DIFF_COLS):
        label = ws.cell(4, col).value
        expected = first_month + i
        if label is None:
            continue
        got = MONTHS.get(str(label).strip().lower())
        if got != expected:
            raise ParseError(
                f"{path} {ws.cell(4, col).coordinate}: forventet måned {expected}, fant {label!r}")
        month_of[col] = f'{year}-{expected:02d}'

    out = []
    for row in range(5, ws.max_row + 1):
        field = ws.cell(row, 2).value
        if field is None or not str(field).strip():
            continue
        field = str(field).strip().upper()
        for col, ym in month_of.items():
            value = _num(ws.cell(row, col).value, f'{path} {ws.cell(row, col).coordinate}')
            if value is not None:
                out.append((field, ym, value))
    return out


def parse_brent(wb, path):
    """Returnerer liste av ('YYYY-MM-DD', pris) for én fil."""
    ws = wb[BRENT_SHEET]
    _quarter(ws, path)  # validerer at arket ser riktig ut
    out = []
    for row in range(5, ws.max_row + 1):
        for dcol, pcol in BRENT_COLS:
            d = ws.cell(row, dcol).value
            if d is None:
                continue
            if not isinstance(d, (datetime, date)):
                raise ParseError(f"{path} {ws.cell(row, dcol).coordinate}: forventet dato, fant {d!r}")
            price = _num(ws.cell(row, pcol).value, f'{path} {ws.cell(row, pcol).coordinate}')
            if price is not None:
                out.append((d.strftime('%Y-%m-%d'), price))
    return out


def load_all(raw_dir=RAW_DIR):
    """Leser alle xlsx-filer. Hovedfiler først, tilleggsfiler til slutt (de vinner)."""
    paths = sorted(glob.glob(os.path.join(raw_dir, '*.xlsx')))
    paths.sort(key=lambda p: (is_supplement(p), os.path.basename(p)))
    if not paths:
        raise ParseError(f"ingen xlsx-filer i {raw_dir}")

    diffs = {}   # (felt, måned) -> (verdi, fil)
    brent = {}   # dato -> (pris, fil)
    conflicts = []
    for path in paths:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
        try:
            for field, ym, value in parse_diffs(wb, path):
                prev = diffs.get((field, ym))
                if prev and prev[0] != value:
                    conflicts.append(('diff', f'{field} {ym}', prev, (value, path)))
                diffs[(field, ym)] = (value, path)
            for d, price in parse_brent(wb, path):
                prev = brent.get(d)
                if prev and prev[0] != price:
                    conflicts.append(('brent', d, prev, (price, path)))
                brent[d] = (price, path)
        finally:
            wb.close()
    return paths, diffs, brent, conflicts


def build_payloads(diffs, brent):
    months = sorted({ym for _, ym in diffs})
    # Fullstendig månedsliste fra første til siste, så grafer kan vise hull.
    y, m = map(int, months[0].split('-'))
    ye, me = map(int, months[-1].split('-'))
    full = []
    while (y, m) <= (ye, me):
        full.append(f'{y}-{m:02d}')
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)

    fields = defaultdict(dict)
    for (field, ym), (value, _) in diffs.items():
        fields[field][ym] = value

    return {
        'normpris_brent.json': {
            'columns': ['date', 'usd'],
            'rows': [[d, brent[d][0]] for d in sorted(brent)],
        },
        'normpris_diff.json': {
            'months': full,
            'fields': {f: dict(sorted(v.items())) for f, v in sorted(fields.items())},
        },
    }


def write_if_changed(name, payload):
    path = os.path.join(OUT_DIR, name)
    text = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(',', ':')) + '\n'
    if os.path.exists(path):
        with open(path, encoding='utf-8') as fh:
            if fh.read() == text:
                print(f"  = {name} uendret")
                return False
    with open(path, 'w', encoding='utf-8') as fh:
        fh.write(text)
    print(f"  + {name} skrevet ({len(text.encode('utf-8')) / 1024:.0f} KB)")
    return True


def main():
    try:
        paths, diffs, brent, conflicts = load_all()
    except ParseError as exc:
        # Avbryt uten å skrive noe, så forrige gyldige datasett består.
        print(f"FEIL: {exc}", file=sys.stderr)
        return 1

    print(f"Leste {len(paths)} filer: {len(diffs)} differensialer, {len(brent)} Brent-dager")
    for kind, key, old, new in conflicts:
        print(f"  ! {kind} {key}: {old[0]} ({os.path.basename(old[1])}) "
              f"overstyres av {new[0]} ({os.path.basename(new[1])})")

    payloads = build_payloads(diffs, brent)
    for name, payload in payloads.items():
        write_if_changed(name, payload)
    return 0


if __name__ == '__main__':
    sys.exit(main())
